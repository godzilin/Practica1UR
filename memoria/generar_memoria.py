"""
Genera los dos documentos de la entrega en esta carpeta:

    memoria.docx    explicación del código (con figuras, máximo 3 páginas)
    preguntas.docx  respuestas a las preguntas del apartado 12

Las figuras salen de figuras.py (ejecutarlo antes si cambian los datos).
Necesita python-docx (pip install python-docx).
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CARPETA = Path(__file__).parent
FIGURAS = CARPETA / "figuras"
# La salida se mantiene junto al script para que la memoria pueda regenerarse
# sin depender de la carpeta desde la que se lance Python.
AUTOR = "Daniel López Farray"
ASIGNATURA = "Redes de Comunicación Industriales"
PRACTICA = "Práctica 1 · Visualizador de un UR3e por TCP/IP"
REPO = "github.com/godzilin/Practica1UR"
AZUL = RGBColor(0x00, 0x5E, 0xB8)          # azul del escudo de Deusto
GRIS = RGBColor(0x55, 0x55, 0x55)
LETRA = "Calibri"


# ---------------------------------------------------------------------
# Formato común
# ---------------------------------------------------------------------
def campo(parrafo, codigo):
    """Campo de Word (PAGE o NUMPAGES): Word pone el número al abrir el documento."""
    # python-docx no ofrece un helper de alto nivel para estos campos; se crea
    # el XML mínimo que Word recalcula al abrir el documento.
    r = parrafo.add_run()
    for tipo in ("begin", "instr", "separate", "texto", "end"):
        if tipo == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = f" {codigo} "
        elif tipo == "texto":
            e = OxmlElement("w:t")
            e.text = "1"
        else:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        r._r.append(e)
    return r


def sin_bordes(tabla):
    # Las tablas se usan como layout auxiliar, no como tablas visibles.
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "nil")
        bordes.append(b)
    tabla._tbl.tblPr.append(bordes)


def linea_inferior(parrafo):
    # Una línea azul separa el encabezado del contenido sin añadir otra imagen.
    bordes = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "005EB8")):
        b.set(qn(k), v)
    bordes.append(b)
    parrafo._p.get_or_add_pPr().append(bordes)


def sombrear(celda, color):
    # El sombreado se aplica al XML de la celda porque es una propiedad de Word.
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(s)


def texto_rico(parrafo, texto, tam=None):
    """Admite **negrita** y `código` dentro del texto."""
    # Se divide solo en los dos marcadores que usa la memoria; el resto conserva
    # el texto literal y evita introducir un parser Markdown completo.
    for trozo in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            r = parrafo.add_run(trozo[2:-2])
            r.bold = True
        elif trozo.startswith("`"):
            r = parrafo.add_run(trozo[1:-1])
            r.font.name = "Consolas"
        else:
            r = parrafo.add_run(trozo)
        if tam:
            r.font.size = Pt(tam)


class Documento:
    """Generador de documentos Word con el formato comun de la entrega.

    Centraliza estilos, encabezado, pie, figuras, tablas y texto enriquecido
    para que ``memoria.docx`` y ``preguntas.docx`` se construyan igual.
    """

    def __init__(self, tam=11):
        # Cada documento comparte márgenes, tipografía, encabezado y pie para
        # que la entrega tenga un formato uniforme.
        self.tam = tam
        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = LETRA, Pt(tam)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), LETRA)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.08

        s = self.doc.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin, s.bottom_margin = Cm(2.6), Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.0)
        s.header_distance, s.footer_distance = Cm(0.8), Cm(0.8)

        # Encabezado: logo a la izquierda; autor, asignatura y práctica a la derecha
        tabla = s.header.add_table(rows=1, cols=2, width=Cm(17))
        sin_bordes(tabla)
        tabla.cell(0, 0).paragraphs[0].add_run().add_picture(str(CARPETA / "logo_deusto.png"),
                                                             height=Cm(1.2))
        p = tabla.cell(0, 1).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(AUTOR + "\n")
        r.bold, r.font.size = True, Pt(10)
        r = p.add_run(f"{ASIGNATURA}\n{PRACTICA}")
        r.font.size, r.font.color.rgb = Pt(8.5), GRIS
        s.header.paragraphs[0].paragraph_format.space_after = Pt(0)
        linea_inferior(s.header.add_paragraph())

        # Pie: Página X de Y
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for t in ("Página ", "PAGE", " de ", "NUMPAGES"):
            r = campo(p, t) if t.isupper() else p.add_run(t)
            r.font.size, r.font.color.rgb = Pt(9), GRIS

    def titulo(self, texto, subtitulo):
        """Añade el titulo principal y su subtitulo."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(texto)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(20), AZUL
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(subtitulo)
        r.font.size, r.font.color.rgb = Pt(11), GRIS

    def seccion(self, texto):
        """Añade un encabezado de seccion que permanece junto al parrafo siguiente."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(texto)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(14), AZUL

    def parrafo(self, texto, tam=None):
        """Añade un parrafo admitiendo negrita y fragmentos de codigo."""
        p = self.doc.add_paragraph()
        texto_rico(p, texto, tam)
        return p

    def punto(self, texto):
        """Añade un elemento de lista con el mismo formato enriquecido."""
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        texto_rico(p, texto)

    def figura(self, nombre, ancho_cm, pie=None):
        """Inserta una figura centrada y, opcionalmente, su pie."""
        # El pie se mantiene junto a la figura para que Word no los separe.
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2 if pie else 6)
        p.paragraph_format.keep_with_next = bool(pie)
        p.add_run().add_picture(str(FIGURAS / nombre), width=Cm(ancho_cm))
        if pie:
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(8)
            r = q.add_run(pie)
            r.italic, r.font.size, r.font.color.rgb = True, Pt(9.5), GRIS

    def lado_a_lado(self, nombre, ancho_cm, puntos, pie=None):
        """Figura a la izquierda y frases cortas a la derecha (ahorra espacio)."""
        t = self.doc.add_table(rows=1, cols=2)
        sin_bordes(t)
        izq, der = t.cell(0, 0), t.cell(0, 1)
        izq.width, der.width = Cm(ancho_cm + 0.3), Cm(17 - ancho_cm - 0.3)
        p = izq.paragraphs[0]
        p.add_run().add_picture(str(FIGURAS / nombre), width=Cm(ancho_cm))
        if pie:
            q = izq.add_paragraph()
            r = q.add_run(pie)
            r.italic, r.font.size, r.font.color.rgb = True, Pt(9), GRIS
        for i, texto in enumerate(puntos):
            q = der.paragraphs[0] if i == 0 else der.add_paragraph()
            q.paragraph_format.space_after = Pt(5)
            texto_rico(q, "• " + texto)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def tabla(self, filas, anchos):
        """Inserta una tabla y aplica un estilo azul a su primera fila."""
        # La primera fila se trata como cabecera; las demás son datos normales.
        t = self.doc.add_table(rows=len(filas), cols=len(filas[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, fila in enumerate(filas):
            for j, texto in enumerate(fila):
                c = t.cell(i, j)
                c.width = Cm(anchos[j])
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                texto_rico(p, texto, self.tam - 1)
                if i == 0:
                    sombrear(c, "005EB8")
                    for r in p.runs:
                        r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def guardar(self, nombre):
        """Escribe el documento en la carpeta de memoria."""
        # Las rutas de salida son relativas a memoria/, nunca al directorio actual.
        self.doc.save(CARPETA / nombre)
        print("Guardado", CARPETA / nombre)


# ---------------------------------------------------------------------
# Memoria
# ---------------------------------------------------------------------
m = Documento(tam=11)
m.titulo("Visualizador de un UR3e por TCP/IP", f"Memoria de la práctica · código en {REPO}")
m.parrafo("Cuatro programas que leen por la red el estado de un brazo **UR3e** (real o URSim) y lo muestran "
          "en 3D con **PyBullet**, de menos a más: una interfaz, otra más rápida, varios robots y, por último, "
          "una aplicación que además da órdenes al robot. Todos se lanzan con F5 y leen las IP de `robots.json`.")

m.seccion("A · Interfaz primaria (primaria.py)")
m.figura("mensaje.png", 13.5, "Fig. 1 · De los bytes que llegan por TCP a los 6 ángulos del robot.")
m.punto("**Un recv() no es un mensaje:** en un minuto hubo 1110 `recv()` de 11 tamaños distintos. "
        "Se acumula en un buffer y se corta leyendo la longitud de cada mensaje.")
m.punto("Del mensaje de estado (tipo 16) solo interesa el **subpaquete 1**; los demás se saltan con su longitud.")
m.punto("Números en **big-endian** (`>`). Si se corta la conexión, **reconecta solo**.")

m.seccion("B · Interfaz de tiempo real (tiempo_real.py)")
m.figura("tiempo_real.png", 13.5, "Fig. 2 · Paquete del puerto 30003 y ritmo de llegada frente al de dibujo.")
m.punto("Estructura fija: los ángulos están **siempre en el byte 252**. No hace falta recorrer subpaquetes.")
m.punto("Llegan ~500 paquetes/s (URSim da ~480), demasiados para dibujarlos todos: se leen todos, **se guarda "
        "solo el último** y se dibuja 30 veces por segundo.")

m.seccion("C · Varios robots (varios_robots.py)")
m.lado_a_lado("captura_varios.png", 8.5, [
    "**Un hilo por robot**, cada uno con su socket y su buffer: si uno se cuelga, solo se bloquea su hilo.",
    "Los hilos dejan el último estado en un diccionario protegido con un **Lock**; un único bucle lo lee y dibuja.",
    "Si un robot deja de enviar 1 s, su nombre pasa a rojo **«SIN DATOS»** y se queda quieto.",
], "Fig. 3 · Dos URSim y un robot apagado.")
m.figura("arquitectura.png", 13.5, "Fig. 4 · Estructura de C y D: recibir, guardar el último estado y dibujar.")

m.seccion("D · Enseñar y reproducir (ensenar_reproducir.py)")
m.lado_a_lado("captura.png", 8.5, [
    "**Enseñar:** Espacio guarda la postura del robot como punto.",
    "**Previsualizar:** un robot «fantasma» azul recorre los puntos sin mover el real.",
    "**Ejecutar:** envía un `movej` por punto por el mismo socket 30001 y **confirma** cada uno.",
], "Fig. 5 · Gris: el robot. Azul: el punto que va a ejecutar.")
m.figura("estados.png", 10, "Fig. 6 · Máquina de estados. Ante un error o parada se detiene y muestra el motivo.")
m.figura("movej.png", 13.5, "Fig. 7 · Datos reales de URSim. El robot no responde a un movej: el punto se da por "
                            "bueno cuando el programa termina y el ángulo está a menos de 0,01 rad.")

m.seccion("Además")
m.punto(f"**Repositorio GitHub** ({REPO}): cada cambio es un commit; README con instalación y uso.")
m.punto("**A prueba de fallos:** reconexión automática, mensajes corruptos descartados, configuración "
        "validada. Probado con 20 000 mensajes rotos, cortes de red y URSim reiniciado.")
m.punto("**Seguridad:** `stopj` si el programa se cierra con el robot en marcha; velocidad ≤ 0,3 rad/s; "
        "una parada de emergencia nunca se desbloquea por red.")

m.seccion("Robot real")
m.parrafo("**[Completar tras la sesión: qué UR3e se usó y diferencias con URSim.]**")
m.guardar("memoria.docx")


# ---------------------------------------------------------------------
# Preguntas del apartado 12
# ---------------------------------------------------------------------
q = Documento(tam=10.5)
q.titulo("Preguntas de comprensión", "Apartado 12 · respuestas con datos medidos en URSim")


def pregunta(n, enunciado, respuesta, figura=None, ancho=None):
    p = q.doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{n}. {enunciado}")
    r.bold, r.font.color.rgb = True, AZUL
    q.parrafo(respuesta)
    if figura:
        q.figura(figura, ancho)


pregunta(1, "¿Coincide cada recv() con un mensaje? ¿Qué consecuencia tiene?",
         "**No.** En un minuto hubo 1110 `recv()` de 11 tamaños distintos (de 9 a 1477 bytes): 1021 trajeron "
         "un mensaje y 89 trajeron varios; también puede llegar uno a medias. Por eso nunca se procesa un "
         "`recv()` directamente: se acumula en un buffer y se corta por la longitud de cada mensaje.",
         "recv.png", 14)
pregunta(2, "¿Por qué el mensaje y cada subpaquete empiezan con su longitud? ¿Y un tipo desconocido?",
         "TCP es un flujo de bytes sin separadores: la longitud dice dónde acaba cada pieza. Un subpaquete de "
         "tipo desconocido se **salta** sumando su longitud, así el programa sigue funcionando aunque una "
         "versión nueva del robot añada datos.")
pregunta(3, "¿Qué pasa si se decodifica q_actual sin big-endian?",
         "Salen valores absurdos porque los bytes se leen al revés. Ejemplo: −1,5708 rad (−90°) viaja como "
         "`BF F9 21 FF 2E 48 E8 A7`; leído en el orden del PC (little-endian) da **−1,9·10⁻¹¹⁶**.")
pregunta(4, "Primaria o tiempo real: ¿cuál para cada uso?",
         "**Primaria:** 10 mensajes/s, poco tráfico, incluye errores y acepta órdenes. **Tiempo real:** "
         "~500 paquetes/s de estructura fija, menos retraso y unas 50 veces más tráfico. Para un **panel de "
         "supervisión**, la primaria. Para **registrar un movimiento rápido**, la de tiempo real (un dato cada 2 ms).")
pregunta(5, "¿Por qué con NAT no se llega a URSim y en modo puente sí? ¿Quién ve la MAC de la VM?",
         "Con NAT la máquina virtual queda escondida tras una red privada de VirtualBox: puede salir, pero desde "
         "el PC no hay una IP a la que conectarse (habría que redirigir puertos). En **modo puente** es un equipo "
         "más de la red del laboratorio, con su propia IP (10.172.18.67). Su MAC la ven el **switch** y los equipos "
         "del mismo segmento, como nuestro PC en su tabla ARP.")
pregunta(6, "Si se corta la red 5 s, ¿qué hace el programa? ¿Y TCP?",
         "**TCP no pierde datos:** el robot los guarda y los reenvía; al volver la red llegan de golpe y "
         "atrasados (en un corte de 1,6 s llegaron 12 estados seguidos). **El programa** espera como máximo "
         "5 s sin datos: si se superan, da la conexión por perdida y reconecta a los 2 s con datos frescos. "
         "Se probó cortando la red con el cortafuegos (para TCP es lo mismo que quitar el cable).")
pregunta(7, "(N) Retraso entre URSim y el robot simulado con cada interfaz.",
         "Medido desde el `movej` hasta ver cambiar el ángulo: **62 ms** con tiempo real y **123 ms** con la "
         "primaria. Unos 60 ms son lo que tarda el robot en arrancar el programa; el resto es el periodo de "
         "publicación (2 ms frente a 100 ms). Se suman hasta 33 ms del dibujo a 30 fotogramas/s. Depende del "
         "periodo de cada interfaz, del ritmo de dibujo, de la red y de la carga del PC y de la VM.",
         "retraso.png", 8)
pregunta(8, "(N) ¿Qué pasaría sin Lock?",
         "El dibujo podría leer el estado mientras el hilo lo escribe y mezclar datos de dos instantes, o "
         "fallar al copiar el diccionario justo cuando otro hilo añade un robot "
         "(`dictionary changed size during iteration`). Son fallos raros e imprevisibles; el Lock los evita.")
pregunta(9, "(S) ¿Por qué no basta con enviar un movej para saber que el robot ha llegado?",
         "Porque el robot no responde: la orden puede rechazarse (sintaxis, modo local), interrumpirse (parada "
         "de protección) o tardar. Se confirma con su estado: `isProgramRunning` pasa a verdadero y luego a "
         "falso, `q_actual` queda a menos de 0,01 rad del objetivo y no llegan errores (mensajes tipo 20) ni "
         "`isProtectiveStopped`. Si no arranca en 2 s, la orden se da por rechazada.")
pregunta(10, "(S) Riesgo de que cualquiera pueda enviar programas al 30001 y dos medidas.",
         "No hay contraseña ni cifrado: cualquier equipo de la red podría mover el robot con personas cerca, "
         "parar la producción o dañar el equipo. **En el robot:** desactivar o restringir en PolyScope los "
         "servicios entrantes (30001–30003) y dejar para supervisión solo interfaces de lectura como la 30011. "
         "**En la red:** poner los robots en una VLAN propia tras un cortafuegos que solo deje conectar a los "
         "puertos de control desde la IP autorizada (zonas y conductos de IEC 62443).")
q.guardar("preguntas.docx")
